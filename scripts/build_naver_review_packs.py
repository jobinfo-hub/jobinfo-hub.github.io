"""Build editable DOCX review packs and a factual PNG summary for Naver drafts.

Requires python-docx and Pillow. Never logs into or posts to Naver.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
DRAFTS = ROOT / "naver-drafts"
OUT = ROOT / "naver-review-packs"
FONT_PATHS = [
    "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]

def font(size):
    for path in FONT_PATHS:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def wrap(draw, value, f, width):
    lines = []
    for chunk in re.split(r"\s+", value.strip()):
        if not chunk:
            continue
        if not lines or draw.textbbox((0, 0), lines[-1] + " " + chunk, font=f)[2] > width:
            lines.append(chunk)
        else:
            lines[-1] += " " + chunk
    return lines

def image_for(draft, target, expired=False, deadline_known=True):
    heading = draft.splitlines()[0].lstrip("# ").strip()
    org = ""
    for line in draft.splitlines():
        if line.startswith("- 모집기관:"):
            org = line.split(":", 1)[1].strip()
            break
    img = Image.new("RGB", (1080, 1080), "#f7f9fc")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((58, 58, 1022, 1022), radius=40, fill="white", outline="#d8e0ed", width=3)
    draw.rounded_rectangle((100, 105, 440, 175), radius=22, fill="#173b73")
    draw.text((125, 117), "CAREERPICK", font=font(37), fill="white")
    status = "마감 공고 · 기록용" if expired else "마감일 확인 필요" if not deadline_known else "채용공고 핵심정리"
    draw.text((108, 240), status, font=font(54), fill="#173b73")
    y = 355
    for line in wrap(draw, heading, font(56), 850)[:6]:
        draw.text((108, y), line, font=font(56), fill="#152438")
        y += 79
    draw.line((108, 835, 970, 835), fill="#d8e0ed", width=3)
    for line in wrap(draw, org, font(36), 840)[:2]:
        draw.text((108, 863), line, font=font(36), fill="#35516c")
        break
    footer = "마감된 공고입니다 · 현재 지원 불가" if expired else "현재 모집 여부를 원문에서 확인하세요"
    draw.text((108, 955), footer, font=font(28), fill="#607389")
    img.save(target)

def make_docx(draft, png, target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(10.5)
    doc.add_picture(str(png), width=Cm(13.8))
    for line in draft.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(target)

def main():
    manifest = json.loads((DRAFTS / "manifest.json").read_text(encoding="utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    current = set()
    for item in manifest:
        pid = item["id"]
        if not re.fullmatch(r"[a-zA-Z0-9_-]+", pid):
            raise ValueError("Unsafe ID")
        draft = (DRAFTS / item["file"]).read_text(encoding="utf-8")
        current.update({f"{pid}-thumbnail.png", f"{pid}.docx"})
        png = OUT / f"{pid}-thumbnail.png"
        image_for(draft, png, expired=item.get("expired", False), deadline_known=item.get("deadline_known", True))
        make_docx(draft, png, OUT / f"{pid}.docx")
    for old in OUT.iterdir():
        if old.is_file() and old.suffix.lower() in (".docx", ".png") and old.name not in current:
            old.unlink()
    (OUT / "README.txt").write_text(
        "커리어픽 검수용 초안입니다. 원문과 조건을 대조한 뒤 사용자가 최종 발행하세요.\n"
        "네이버 블로그 임시저장 및 자동 발행 기능은 포함되어 있지 않습니다.\n",
        encoding="utf-8",
    )
    print(f"Created {len(manifest)} DOCX and thumbnail pairs")

if __name__ == "__main__":
    main()
